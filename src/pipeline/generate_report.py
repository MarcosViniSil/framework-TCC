from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from common.details import ProjectDetails
from domain.csv_model import CSVModel


def generate_excel(
    header: list[str],
    csv_data: list[CSVModel],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = ProjectDetails.get_project_name()

    header_row = 5

    MIN_COLUMN_WIDTH = 15
    MAX_COLUMN_WIDTH = 40
    EXTRA_SPACE = 4

    HEADER_COLOR = "1F4E78"
    HEADER_TEXT_COLOR = "FFFFFF"
    ALTERNATE_ROW_COLOR = "F3F6FA"
    BORDER_COLOR = "D9E1F2"
    TITLE_COLOR = "1F4E78"
    TEXT_COLOR = "333333"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_COLOR,
    )

    alternate_row_fill = PatternFill(
        fill_type="solid",
        fgColor=ALTERNATE_ROW_COLOR,
    )

    header_font = Font(
        bold=True,
        color=HEADER_TEXT_COLOR,
        size=11,
    )

    title_font = Font(
        bold=True,
        color=TITLE_COLOR,
        size=18,
    )

    label_font = Font(
        bold=True,
        color=TEXT_COLOR,
        size=10,
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color=BORDER_COLOR,
        ),
        right=Side(
            style="thin",
            color=BORDER_COLOR,
        ),
        top=Side(
            style="thin",
            color=BORDER_COLOR,
        ),
        bottom=Side(
            style="thin",
            color=BORDER_COLOR,
        ),
    )

    project_name = ProjectDetails.get_project_name()
    project_description = ProjectDetails.get_project_description()

    sheet["A1"] = project_name
    sheet["A1"].font = title_font
    sheet["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    last_column = get_column_letter(len(header))
    sheet.merge_cells(f"A1:{last_column}1")

    sheet["A2"] = project_description
    sheet["A2"].font = label_font
    sheet["A2"].alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    sheet.merge_cells(f"A2:{last_column}2")

    sheet["A4"] = "Generated in"
    sheet["A4"].font = label_font

    sheet["B4"] = datetime.now().strftime("%d/%m/%Y %H:%M")

    sheet["B4"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    for col, value in enumerate(header, start=1):
        cell = sheet.cell(
            row=header_row,
            column=col,
            value=value,
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.row_dimensions[header_row].height = 40

    for row, item in enumerate(
        csv_data,
        start=header_row + 1,
    ):
        values = asdict(item)

        for col, field in enumerate(header, start=1):
            cell = sheet.cell(
                row=row,
                column=col,
                value=values.get(field),
            )

            cell.border = thin_border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            if row % 2 == 0:
                cell.fill = alternate_row_fill

    sheet.freeze_panes = "A6"

    sheet.auto_filter.ref = f"A{header_row}:" f"{last_column}" f"{sheet.max_row}"

    for column in range(1, len(header) + 1):
        column_letter = get_column_letter(column)

        max_length = len(
            str(
                sheet.cell(
                    header_row,
                    column,
                ).value
                or ""
            )
        )

        for row in range(
            header_row + 1,
            sheet.max_row + 1,
        ):
            value = sheet.cell(
                row,
                column,
            ).value

            if value is not None:
                max_length = max(
                    max_length,
                    len(str(value)),
                )

        width = max_length + EXTRA_SPACE

        width = max(
            width,
            MIN_COLUMN_WIDTH,
        )

        width = min(
            width,
            MAX_COLUMN_WIDTH,
        )

        sheet.column_dimensions[column_letter].width = width

    for row in range(
        header_row + 1,
        sheet.max_row + 1,
    ):
        sheet.row_dimensions[row].height = 25

    sheet.sheet_view.zoomScale = 85

    dir_path = Path(ProjectDetails.get_project_default_result_folder())

    dir_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    filename = f"{datetime.now().strftime('%d-%m-%Y-%H-%M')}" "_result.xlsx"

    workbook.save(dir_path / filename)
